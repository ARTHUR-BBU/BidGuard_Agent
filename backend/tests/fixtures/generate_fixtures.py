from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

FIXTURE_DIR = Path(__file__).parent


def _pdf_string(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_pdf() -> None:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    pages = [
        [
            "1 Tender Requirements",
            "1.1 The bidder shall provide a complete technical response.",
            "The response must identify the delivery team and implementation scope.",
        ],
        [
            "2 Commercial Requirements",
            "2.1 The proposal shall state the total price and validity period.",
            "All mandatory forms must be included before the submission deadline.",
        ],
    ]
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        commands = ["BT", "/F1 12 Tf", "72 720 Td", "16 TL"]
        for index, line in enumerate(lines):
            if index:
                commands.append("T*")
            commands.append(f"({_pdf_string(line)}) Tj")
        commands.append("ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(commands).encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    with (FIXTURE_DIR / "sample-tender.pdf").open("wb") as output:
        writer.write(output)


def build_docx() -> None:
    document = Document()
    document.add_heading("1 技术要求", level=1)
    document.add_paragraph("投标文件应说明总体技术方案与实施方法。")
    document.add_heading("1.1 实施范围", level=2)
    document.add_paragraph("1.1 实施范围包括系统部署、培训和验收支持。")
    document.add_heading("2 商务要求", level=1)
    document.add_paragraph("投标报价应包含税费，并说明报价有效期。")
    document.save(str(FIXTURE_DIR / "sample-proposal.docx"))


if __name__ == "__main__":
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    build_pdf()
    build_docx()
