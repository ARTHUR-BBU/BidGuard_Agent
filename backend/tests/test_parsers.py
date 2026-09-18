from __future__ import annotations

import threading
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as WordDocument
from pypdf import PdfWriter
from pypdf.errors import LimitReachedError
from sqlalchemy import func, select

from app.db import Base, GuardedSession, build_engine
from app.documents import docx as docx_parser
from app.documents import pdf as pdf_parser
from app.documents.parser import DocumentParseError, chunk_text, parse_document
from app.documents.storage import sha256_bytes
from app.domain.enums import DocumentRole
from app.domain.schemas import ParseCoverageIssue, ParsedChunk, ParsedDocument
from app.persistence.models import (
    BidProject,
    Document,
    DocumentChunk,
    DocumentVersion,
)
from app.services import ingestion
from app.services.ingestion import ingest_project_document, parse_document_version

FIXTURES = Path(__file__).parent / "fixtures"
SAMPLE_PDF = FIXTURES / "sample-tender.pdf"
SAMPLE_DOCX = FIXTURES / "sample-proposal.docx"


def test_parser_contract_rejects_empty_chunk_text() -> None:
    with pytest.raises(ValueError):
        ParsedChunk(chunk_index=0, page_number=1, section_path=None, text="")

    parsed = ParsedDocument(
        chunks=[
            ParsedChunk(
                chunk_index=0,
                page_number=1,
                section_path=None,
                text="1.1 保留章节编号",
            )
        ],
        total_pages=1,
        parsed_pages=[1],
    )
    assert parsed.chunks[0].text.startswith("1.1")


def test_parser_contract_rejects_impossible_section_ordinals() -> None:
    with pytest.raises(ValueError, match="parsed_sections cannot exceed"):
        ParsedDocument(chunks=[], total_sections=1, parsed_sections=[2])

    chunk = ParsedChunk(
        chunk_index=0,
        page_number=None,
        section_path="资格要求",
        section_ordinal=2,
        text="正文",
    )
    with pytest.raises(ValueError, match="must belong to parsed_sections"):
        ParsedDocument(
            chunks=[chunk],
            total_sections=2,
            parsed_sections=[1],
        )
    with pytest.raises(ValueError, match="require total_sections"):
        ParsedDocument(chunks=[chunk], parsed_sections=[2])

    issue = ParseCoverageIssue(
        code="section_failed",
        section_path="资格要求",
        section_ordinal=2,
    )
    with pytest.raises(ValueError, match="issues require total_sections"):
        ParsedDocument(chunks=[], coverage_issues=[issue])
    with pytest.raises(ValueError, match="cannot exceed total_sections"):
        ParsedDocument(
            chunks=[],
            total_sections=1,
            coverage_issues=[issue],
        )

    unresolved_issue = ParsedDocument(
        chunks=[],
        total_sections=2,
        parsed_sections=[],
        coverage_issues=[issue],
    )
    assert unresolved_issue.coverage_issues[0].section_ordinal == 2


def test_docx_body_without_headings_is_one_physical_section(tmp_path: Path) -> None:
    path = tmp_path / "body-only.docx"
    document = WordDocument()
    document.add_paragraph("没有标题但仍然可定位的正文。")
    document.save(str(path))

    result = parse_document(path)

    assert result.total_sections == 1
    assert result.parsed_sections == [1]
    assert len(result.chunks) == 1
    assert result.chunks[0].section_path is None
    assert result.chunks[0].section_ordinal == 1


def test_real_pdf_fixture_preserves_page_numbers_and_text() -> None:
    result = parse_document(SAMPLE_PDF)

    assert result.total_pages == 2
    assert result.parsed_pages == [1, 2]
    assert result.failed_pages == []
    assert result.ocr_pages == []
    assert result.needs_ocr is False
    assert [chunk.chunk_index for chunk in result.chunks] == list(
        range(len(result.chunks))
    )
    assert {chunk.page_number for chunk in result.chunks} == {1, 2}
    assert all(chunk.section_path is None for chunk in result.chunks)
    assert "1.1" in " ".join(chunk.text for chunk in result.chunks)


def test_real_docx_fixture_preserves_complete_heading_paths() -> None:
    result = parse_document(SAMPLE_DOCX)

    paths = {chunk.section_path for chunk in result.chunks}
    assert "1 技术要求" in paths
    assert "1 技术要求 > 1.1 实施范围" in paths
    assert "2 商务要求" in paths
    implementation = next(
        chunk
        for chunk in result.chunks
        if chunk.section_path == "1 技术要求 > 1.1 实施范围"
    )
    assert "1.1" in implementation.text
    assert result.failed_pages == []
    assert result.needs_ocr is False
    assert result.total_sections == 3
    assert result.parsed_sections == [1, 2, 3]
    assert [chunk.section_ordinal for chunk in result.chunks] == [1, 2, 3]


class _FakePage:
    def __init__(
        self,
        text: str = "",
        *,
        raises: bool = False,
        has_image: bool = False,
    ) -> None:
        self._text = text
        self._raises = raises
        self.images = [object()] if has_image else []

    def extract_text(self) -> str:
        if self._raises:
            raise RuntimeError("parser internals must not escape")
        return self._text


class _FakeReader:
    def __init__(self, _path: Path, pages: list[_FakePage]) -> None:
        self.pages = pages
        self.attachment_list: list[object] = []


def test_pdf_omits_blank_pages_marks_low_text_and_records_page_failure(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    pages = [
        _FakePage(),
        _FakePage("页眉文字", has_image=True),
        _FakePage("3.1 " + "可核查正文" * 20),
        _FakePage(raises=True),
    ]
    monkeypatch.setattr(
        pdf_parser,
        "PdfReader",
        lambda path, **_kwargs: _FakeReader(path, pages),
    )
    path = tmp_path / "coverage.pdf"
    path.write_bytes(b"%PDF-real-reader-is-monkeypatched")

    result = pdf_parser.parse_pdf(path)

    assert result.total_pages == 4
    assert result.parsed_pages == [3]
    assert result.blank_pages == [1]
    assert result.failed_pages == [4]
    assert result.ocr_pages == [2]
    assert result.needs_ocr is True
    assert [chunk.page_number for chunk in result.chunks] == [3]
    assert all(chunk.text.strip() for chunk in result.chunks)
    assert any(issue.code == "image_not_extracted" for issue in result.coverage_issues)
    assert any(issue.page_number == 2 for issue in result.coverage_issues)
    assert any(issue.code == "blank_page" for issue in result.coverage_issues)


def test_docx_heading_jump_does_not_invent_missing_headings(tmp_path: Path) -> None:
    path = tmp_path / "heading-jump.docx"
    document = WordDocument()
    document.add_heading("1 一级标题", level=1)
    document.add_heading("1.1.1 三级标题", level=3)
    document.add_paragraph("真实正文")
    document.save(str(path))

    result = parse_document(path)

    paths = {chunk.section_path for chunk in result.chunks}
    assert "1 一级标题 > 1.1.1 三级标题" in paths
    assert all("untitled" not in (path or "").lower() for path in paths)


def test_docx_records_table_and_embedded_visual_coverage_limits(
    tmp_path: Path,
) -> None:
    path = tmp_path / "complex.docx"
    document = WordDocument()
    document.add_heading("1 资格要求", level=1)
    document.add_paragraph("投标人应提交完整材料。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "证书"
    table.cell(0, 1).text = "有效期"
    table.cell(1, 0).text = "ISO 9001"
    table.cell(1, 1).text = "2030-12-31"
    document.save(str(path))

    result = parse_document(path)

    assert any(issue.code == "table_structure_not_preserved" for issue in result.coverage_issues)
    assert any(issue.section_path == "1 资格要求" for issue in result.coverage_issues)
    assert "ISO 9001" in " ".join(chunk.text for chunk in result.chunks)


def test_chunking_has_bounded_overlap_and_never_stalls_or_duplicates() -> None:
    source = "".join(chr(0x4E00 + index) for index in range(6_401))

    chunks = chunk_text(source, chunk_size=3_000, overlap=300)

    assert [len(chunk) for chunk in chunks] == [3_000, 3_000, 1_001]
    assert chunks[0][-300:] == chunks[1][:300]
    assert chunks[1][-300:] == chunks[2][:300]
    assert all(chunks)
    assert len(set(chunks)) == len(chunks)
    with pytest.raises(ValueError):
        chunk_text(source, chunk_size=300, overlap=300)


def test_ingestion_parses_each_version_without_replacing_older_chunks(
    db_session, tmp_path: Path
) -> None:
    project = BidProject(name="版本解析")
    db_session.add(project)
    db_session.commit()

    pdf_bytes = SAMPLE_PDF.read_bytes()
    first = ingest_project_document(
        db_session,
        tmp_path / "uploads",
        project.id,
        DocumentRole.TENDER,
        SAMPLE_PDF.name,
        BytesIO(pdf_bytes),
        digest=sha256_bytes(pdf_bytes),
        size=len(pdf_bytes),
    )
    first_chunk_ids = list(
        db_session.scalars(
            select(DocumentChunk.id).where(
                DocumentChunk.document_version_id == first.version.id
            )
        )
    )
    assert first.version.parse_status == "partial_failure"
    assert first.version.parse_error_code == "incomplete_coverage"
    assert first_chunk_ids

    docx_bytes = SAMPLE_DOCX.read_bytes()
    second = ingest_project_document(
        db_session,
        tmp_path / "uploads",
        project.id,
        DocumentRole.TENDER,
        SAMPLE_DOCX.name,
        BytesIO(docx_bytes),
        digest=sha256_bytes(docx_bytes),
        size=len(docx_bytes),
    )
    assert second.version.version_number == 2
    assert second.version.parse_status == "parsed"
    assert second.version.chunks

    parse_document_version(db_session, second.version.id, tmp_path / "uploads")
    db_session.expire_all()
    assert list(
        db_session.scalars(
            select(DocumentChunk.id).where(
                DocumentChunk.document_version_id == first.version.id
            )
        )
    ) == first_chunk_ids
    assert (
        db_session.scalar(
            select(func.count()).select_from(DocumentChunk).where(
                DocumentChunk.document_version_id == second.version.id
            )
        )
        == len(second.version.chunks)
    )


def test_ingestion_keeps_upload_and_records_stable_safe_parse_failure(
    db_session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    project = BidProject(name="失败恢复")
    db_session.add(project)
    db_session.commit()
    secret = "C:/private/customer-name.pdf: super secret parser traceback"

    def fail_parse(_path: Path) -> ParsedDocument:
        raise RuntimeError(secret)

    monkeypatch.setattr(ingestion, "parse_document", fail_parse)
    content = b"%PDF-corrupt-but-stored"
    result = ingest_project_document(
        db_session,
        tmp_path / "uploads",
        project.id,
        DocumentRole.TENDER,
        "broken.pdf",
        BytesIO(content),
        digest=sha256_bytes(content),
        size=len(content),
    )

    assert Path(result.version.storage_path).read_bytes() == content
    assert result.version.parse_status == "failed"
    assert result.version.parse_error_code == "parse_failed"
    assert result.version.parse_error == "Document parsing failed"
    assert secret not in result.version.parse_error
    assert result.version.chunks == []


def test_failed_reparse_keeps_last_successful_snapshot(
    db_session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    project = BidProject(name="安全重解析")
    db_session.add(project)
    db_session.commit()
    content = SAMPLE_PDF.read_bytes()
    result = ingest_project_document(
        db_session,
        tmp_path / "uploads",
        project.id,
        DocumentRole.TENDER,
        SAMPLE_PDF.name,
        BytesIO(content),
        digest=sha256_bytes(content),
        size=len(content),
    )
    original_chunks = [(chunk.chunk_index, chunk.text) for chunk in result.version.chunks]
    original_coverage = result.version.parse_coverage

    def fail_parse(_path: Path) -> ParsedDocument:
        raise RuntimeError("transient parser failure")

    monkeypatch.setattr(ingestion, "parse_document", fail_parse)
    assert not parse_document_version(
        db_session, result.version.id, tmp_path / "uploads"
    )
    db_session.refresh(result.version)

    assert result.version.parse_status == "partial_failure"
    assert result.version.parse_error_code == "parse_failed"
    assert result.version.parse_error == (
        "Latest parse attempt failed; previous parsed content retained"
    )
    assert result.version.parse_coverage == original_coverage
    assert [(chunk.chunk_index, chunk.text) for chunk in result.version.chunks] == original_chunks


def test_repeat_upload_retries_failed_but_does_not_reparse_completed_version(
    db_session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    project = BidProject(name="重复上传")
    db_session.add(project)
    db_session.commit()
    storage_root = tmp_path / "uploads"
    content = SAMPLE_PDF.read_bytes()
    digest = sha256_bytes(content)
    first = ingest_project_document(
        db_session,
        storage_root,
        project.id,
        DocumentRole.TENDER,
        SAMPLE_PDF.name,
        BytesIO(content),
        digest=digest,
        size=len(content),
    )
    assert first.version.parse_status == "partial_failure"
    original = ingestion.parse_document_version
    calls: list[int] = []

    def recording_parse(session, version_id: int, root: Path) -> bool:
        calls.append(version_id)
        return original(session, version_id, root)

    monkeypatch.setattr(ingestion, "parse_document_version", recording_parse)
    repeated = ingest_project_document(
        db_session,
        storage_root,
        project.id,
        DocumentRole.TENDER,
        SAMPLE_PDF.name,
        BytesIO(content),
        digest=digest,
        size=len(content),
    )
    assert repeated.created is False
    assert calls == []

    repeated.version.parse_status = "failed"
    db_session.commit()
    retried = ingest_project_document(
        db_session,
        storage_root,
        project.id,
        DocumentRole.TENDER,
        SAMPLE_PDF.name,
        BytesIO(content),
        digest=digest,
        size=len(content),
    )
    assert retried.created is False
    assert calls == [first.version.id]
    assert retried.version.parse_status == "partial_failure"


def test_source_path_size_and_digest_are_verified_before_parsing(
    db_session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    project = BidProject(name="来源校验")
    document = Document(
        project=project,
        role="tender",
        display_name="outside.pdf",
    )
    outside = tmp_path / "outside.pdf"
    outside.write_bytes(SAMPLE_PDF.read_bytes())
    version = DocumentVersion(
        document=document,
        version_number=1,
        sha256=sha256_bytes(outside.read_bytes()),
        size_bytes=outside.stat().st_size,
        storage_path=str(outside),
        parse_status="pending",
    )
    db_session.add(version)
    db_session.commit()
    called = False

    def should_not_parse(_path: Path) -> ParsedDocument:
        nonlocal called
        called = True
        raise AssertionError("integrity gate was bypassed")

    monkeypatch.setattr(ingestion, "parse_document", should_not_parse)
    parse_document_version(db_session, version.id, tmp_path / "uploads")

    assert called is False
    assert version.parse_status == "failed"
    assert version.parse_error_code == "source_integrity_failed"
    assert version.parse_coverage is not None
    assert version.parse_coverage["coverage_issues"][0]["code"] == "source_unverified"
    assert outside.exists()


def test_docx_zip_resource_limit_fails_closed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(docx_parser, "MAX_DOCX_UNCOMPRESSED_BYTES", 100)

    with pytest.raises(DocumentParseError) as caught:
        parse_document(SAMPLE_DOCX)

    assert caught.value.code == "resource_limit_exceeded"


def test_pdf_page_limit_fails_closed(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    pages = [_FakePage("足够长的正文" * 10), _FakePage("足够长的正文" * 10)]
    monkeypatch.setattr(pdf_parser, "MAX_PDF_PAGES", 1)
    monkeypatch.setattr(
        pdf_parser,
        "PdfReader",
        lambda path, **_kwargs: _FakeReader(path, pages),
    )
    path = tmp_path / "too-many.pdf"
    path.write_bytes(b"%PDF-reader-monkeypatched")

    with pytest.raises(DocumentParseError) as caught:
        pdf_parser.parse_pdf(path)

    assert caught.value.code == "resource_limit_exceeded"


def test_newer_parse_attempt_wins_and_chunks_are_not_mixed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    engine = build_engine(f"sqlite:///{tmp_path / 'concurrent.db'}")
    Base.metadata.create_all(engine)
    storage_root = tmp_path / "uploads"
    content = SAMPLE_PDF.read_bytes()
    digest = sha256_bytes(content)
    stored_path = storage_root / digest[:2] / f"{digest}.pdf"
    stored_path.parent.mkdir(parents=True)
    stored_path.write_bytes(content)
    with GuardedSession(engine, expire_on_commit=False) as seed:
        version = DocumentVersion(
            document=Document(
                project=BidProject(name="并发解析"),
                role="tender",
                display_name="sample.pdf",
            ),
            version_number=1,
            sha256=digest,
            size_bytes=len(content),
            storage_path=str(stored_path),
            parse_status="pending",
        )
        seed.add(version)
        seed.commit()
        version_id = version.id

    first_entered = threading.Event()
    release_first = threading.Event()
    counter_lock = threading.Lock()
    call_number = 0

    def controlled_parse(_path: Path) -> ParsedDocument:
        nonlocal call_number
        with counter_lock:
            call_number += 1
            current = call_number
        if current == 1:
            first_entered.set()
            assert release_first.wait(10)
            text = "older result"
        else:
            text = "newer result"
        return ParsedDocument(
            chunks=[
                ParsedChunk(
                    chunk_index=0,
                    page_number=1,
                    section_path=None,
                    text=text,
                )
            ],
            total_pages=1,
            parsed_pages=[1],
        )

    monkeypatch.setattr(ingestion, "parse_document", controlled_parse)
    first_result: list[bool] = []

    def run_first() -> None:
        with GuardedSession(engine, expire_on_commit=False) as session:
            first_result.append(
                parse_document_version(session, version_id, storage_root)
            )

    first = threading.Thread(target=run_first)
    first.start()
    assert first_entered.wait(10)
    with GuardedSession(engine, expire_on_commit=False) as second_session:
        assert parse_document_version(second_session, version_id, storage_root)
    release_first.set()
    first.join(10)
    assert not first.is_alive()

    with GuardedSession(engine) as check:
        persisted = check.get(DocumentVersion, version_id)
        assert persisted is not None
        assert persisted.parse_status == "parsed"
        assert [chunk.text for chunk in persisted.chunks] == ["newer result"]
    assert first_result == [False]
    engine.dispose()


def test_parsed_document_coverage_issue_requires_a_location() -> None:
    with pytest.raises(ValueError):
        ParseCoverageIssue(code="image_not_extracted")


def test_pdf_embedded_attachment_and_table_capability_are_explicit(
    tmp_path: Path,
) -> None:
    path = tmp_path / "attachment.pdf"
    writer = PdfWriter()
    writer.append(str(SAMPLE_PDF))
    writer.add_attachment("requirements.xlsx", b"not parsed")
    with path.open("wb") as output:
        writer.write(output)

    result = parse_document(path)

    codes = [issue.code for issue in result.coverage_issues]
    assert "embedded_attachment_not_extracted" in codes
    assert "pdf_table_detection_unavailable" in codes


def test_docx_non_body_story_content_is_not_silently_ignored(tmp_path: Path) -> None:
    path = tmp_path / "header.docx"
    document = WordDocument()
    document.sections[0].header.paragraphs[0].text = "页眉中的项目编号"
    document.sections[0].footer.paragraphs[0].text = "页脚中的保密说明"
    document.add_heading("1 正文", level=1)
    document.add_paragraph("可解析的正文内容。")
    document.save(str(path))

    result = parse_document(path)

    story_issues = [
        issue
        for issue in result.coverage_issues
        if issue.code == "non_body_story_not_extracted"
    ]
    assert len(story_issues) == 2
    assert all(issue.section_path and "word/" in issue.section_path for issue in story_issues)
    assert result.total_sections == 1
    assert result.parsed_sections == [1]


def test_all_pdf_page_extraction_failures_are_incomplete_not_empty(
    db_session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        ingestion,
        "parse_document",
        lambda _path: ParsedDocument(
            chunks=[],
            total_pages=2,
            failed_pages=[1, 2],
            coverage_issues=[
                ParseCoverageIssue(code="page_extraction_failed", page_number=1),
                ParseCoverageIssue(code="page_extraction_failed", page_number=2),
            ],
        ),
    )
    project = BidProject(name="全页失败")
    db_session.add(project)
    db_session.commit()
    content = SAMPLE_PDF.read_bytes()

    result = ingest_project_document(
        db_session,
        tmp_path / "uploads",
        project.id,
        DocumentRole.TENDER,
        SAMPLE_PDF.name,
        BytesIO(content),
        digest=sha256_bytes(content),
        size=len(content),
    )

    assert result.version.parse_status == "failed"
    assert result.version.parse_error_code == "incomplete_coverage"


def test_pdf_input_limit_is_checked_before_reader_construction(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "oversized.pdf"
    with path.open("wb") as output:
        output.truncate(11)
    monkeypatch.setattr(pdf_parser, "MAX_PDF_INPUT_BYTES", 10)
    called = False

    def reader_must_not_run(*_args, **_kwargs):
        nonlocal called
        called = True
        raise AssertionError("reader constructed before the input limit gate")

    monkeypatch.setattr(pdf_parser, "PdfReader", reader_must_not_run)

    with pytest.raises(DocumentParseError) as caught:
        parse_document(path)
    assert caught.value.code == "resource_limit_exceeded"
    assert called is False


def test_pdf_decompression_limit_is_not_downgraded_to_page_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    class LimitedPage(_FakePage):
        def extract_text(self) -> str:
            raise LimitReachedError("decompression output limit reached")

    monkeypatch.setattr(
        pdf_parser,
        "PdfReader",
        lambda path, **_kwargs: _FakeReader(path, [LimitedPage()]),
    )
    path = tmp_path / "bomb.pdf"
    path.write_bytes(b"%PDF-reader-monkeypatched")

    with pytest.raises(DocumentParseError) as caught:
        parse_document(path)
    assert caught.value.code == "resource_limit_exceeded"


def test_stale_parsing_lease_is_reclaimed_but_fresh_lease_is_not(
    db_session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    project = BidProject(name="解析租约")
    db_session.add(project)
    db_session.commit()
    storage_root = tmp_path / "uploads"
    content = SAMPLE_PDF.read_bytes()
    digest = sha256_bytes(content)
    first = ingest_project_document(
        db_session,
        storage_root,
        project.id,
        DocumentRole.TENDER,
        SAMPLE_PDF.name,
        BytesIO(content),
        digest=digest,
        size=len(content),
    )
    calls: list[int] = []
    original = ingestion.parse_document_version

    def recording_parse(session, version_id: int, root: Path) -> bool:
        calls.append(version_id)
        return original(session, version_id, root)

    monkeypatch.setattr(ingestion, "parse_document_version", recording_parse)
    first.version.parse_status = "parsing"
    first.version.parse_attempt_started_at = datetime.now(UTC)
    db_session.commit()
    ingest_project_document(
        db_session,
        storage_root,
        project.id,
        DocumentRole.TENDER,
        SAMPLE_PDF.name,
        BytesIO(content),
        digest=digest,
        size=len(content),
    )
    assert calls == []

    first.version.parse_attempt_started_at = datetime.now(UTC) - timedelta(minutes=16)
    db_session.commit()
    ingest_project_document(
        db_session,
        storage_root,
        project.id,
        DocumentRole.TENDER,
        SAMPLE_PDF.name,
        BytesIO(content),
        digest=digest,
        size=len(content),
    )
    assert calls == [first.version.id]


def test_repeated_same_heading_creates_distinct_physical_sections(
    tmp_path: Path,
) -> None:
    path = tmp_path / "repeated-headings.docx"
    document = WordDocument()
    document.add_heading("资格要求", level=1)
    document.add_paragraph("第一处要求提供营业执照。")
    document.add_heading("资格要求", level=1)
    document.add_paragraph("第二处要求提供审计报告。")
    document.save(str(path))

    result = parse_document(path)

    assert result.total_sections == 2
    assert result.parsed_sections == [1, 2]
    assert [chunk.section_path for chunk in result.chunks] == [
        "资格要求",
        "资格要求",
    ]
    assert [chunk.section_ordinal for chunk in result.chunks] == [1, 2]
    assert "营业执照" in result.chunks[0].text
    assert "审计报告" in result.chunks[1].text
